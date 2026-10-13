#!/usr/bin/env python3
"""slides.pptx — the deck as PowerPoint, for a venue that wants .pptx.

WHY IMAGES AND NOT TEXT BOXES. Converting LaTeX layout into native PowerPoint
shapes means PowerPoint re-flows it: fonts substitute, line breaks move, and
a slide that was checked for orphans on this machine is not the slide that
appears on theirs. One full-bleed image per slide cannot re-flow. It is not
editable, which for a deck that gets shown and not edited is the right trade.

Speaker notes are carried across into the PowerPoint notes field, so presenter
view works on their laptop.

Needs: python-pptx, pdftoppm (poppler).
"""
import re, subprocess, sys, tempfile, pathlib
from pptx import Presentation
from pptx.util import Inches, Emu

HERE = pathlib.Path(__file__).parent
PDF  = HERE / "slides.pdf"
OUT  = HERE / "slides.pptx"
DPI  = 150
W_IN, H_IN = 13.3333, 7.5          # PowerPoint's own 16:9, not LaTeX's

def notes_from_body():
    """Pull each \\note{...} out of body.tex and flatten it to plain text."""
    s = (HERE / "body.tex").read_text()
    spans, i = [], 0
    while True:
        i = s.find("\\note{", i)
        if i < 0: break
        j, depth = i + 6, 1
        while depth:
            if s[j] == "{": depth += 1
            elif s[j] == "}": depth -= 1
            j += 1
        spans.append(s[i+6:j-1]); i = j
    out = []
    for n in spans:
        t = n
        t = re.sub(r'\\begin\{itemize\}\\itemsep1pt', '', t)
        t = re.sub(r'\\(?:begin|end)\{itemize\}', '', t)
        t = t.replace("\\item", "\u2022")
        t = re.sub(r'\\textbf\{([^{}]*)\}', r'\1', t)
        t = re.sub(r'\\textit\{([^{}]*)\}', r'\1', t)
        t = t.replace("\\quad", "  ").replace("\\\\", "\n")
        t = t.replace("$\\leftarrow$", "<-").replace("---", "\u2014")
        t = t.replace("``", '"').replace("''", '"')
        t = re.sub(r'\\[a-zA-Z]+', '', t)
        t = re.sub(r'[ \t]+', ' ', t)
        t = re.sub(r'\n\s*\n+', '\n', t).strip()
        out.append(t)
    return out

def main():
    if not PDF.exists(): sys.exit("build slides.pdf first (make)")
    notes = notes_from_body()
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["pdftoppm", "-jpeg", "-r", str(DPI), "-jpegopt", "quality=88",
                        str(PDF), f"{td}/s"], check=True)
        pages = sorted(pathlib.Path(td).glob("s-*.jpg"))
        if not pages: sys.exit("pdftoppm produced nothing")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
        blank = prs.slide_layouts[6]           # 6 is the blank layout
        for i, img in enumerate(pages):
            sl = prs.slides.add_slide(blank)
            sl.shapes.add_picture(str(img), Emu(0), Emu(0),
                                  width=prs.slide_width, height=prs.slide_height)
            if i < len(notes) and notes[i]:
                sl.notes_slide.notes_text_frame.text = notes[i]
        prs.save(OUT)
    mb = OUT.stat().st_size / 1e6
    print(f"slides.pptx  {len(pages)} slides, {mb:.1f} MB, {W_IN}x{H_IN} in, "
          f"notes on {sum(1 for n in notes if n)} of them")

if __name__ == "__main__":
    main()
